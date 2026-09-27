"""Lê as 3 planilhas do mês e gera o JSON único usado pelo dashboard e pelo relatório.

Uso: python3 scripts/extrair_dados.py <pasta_do_mes> <saida.json>
"""
import json
import sys
from datetime import datetime
from pathlib import Path

import openpyxl

# Categorias de reembolso que têm item correspondente na EAP (vínculo sugerido,
# usado só para detectar estouro de item; o realizado oficial não muda).
MAPA_REEMBOLSO_EAP = {"Taxas e ART": "001.01.01.01"}


def achar(pasta, prefixo):
    arqs = sorted(Path(pasta).glob(f"{prefixo}*.xlsx"))
    if not arqs:
        sys.exit(f"Arquivo '{prefixo}*.xlsx' não encontrado em {pasta}")
    return arqs[0]


def linhas_tabela(ws, primeira_coluna):
    """Devolve dicts a partir da linha de cabeçalho cuja 1ª célula preenchida é `primeira_coluna`.
    Para na linha TOTAL ou na primeira linha vazia."""
    rows = list(ws.iter_rows(values_only=True))
    for i, r in enumerate(rows):
        cel = [c for c in r if c is not None]
        if cel and cel[0] == primeira_coluna:
            idx = {c: j for j, c in enumerate(r) if c is not None}
            out = []
            for r2 in rows[i + 1:]:
                vals = [c for c in r2 if c is not None]
                if not vals or vals[0] == "TOTAL":
                    break
                out.append({k: r2[j] for k, j in idx.items()})
            return out
    sys.exit(f"Cabeçalho '{primeira_coluna}' não encontrado na aba {ws.title}")


def iso(d):
    return d.isoformat(timespec="minutes") if isinstance(d, datetime) else d


def main(pasta, saida):
    fin = openpyxl.load_workbook(achar(pasta, "2_Controle_Financeiro"), data_only=True, read_only=True)
    reemb = openpyxl.load_workbook(achar(pasta, "1_Reembolso"), data_only=True, read_only=True)
    obra = openpyxl.load_workbook(achar(pasta, "3_Acompanhamento_Obra"), data_only=True, read_only=True)

    eap = [{
        "cod": str(r["Código"]), "nivel": int(r["Nível"]), "desc": r["Descrição"].strip(),
        "orc": float(r["Orçado (R$)"] or 0), "real_planilha": float(r["Realizado (R$)"] or 0),
    } for r in linhas_tabela(fin["Orçado x Realizado"], "Código")]

    lanc = [{
        "n": r["Nº"], "data": iso(r["Data e hora"]), "centro": r["Centro de custo"],
        "cod": str(r["Código EAP"]), "desc": r["Descrição"], "favorecido": r["Favorecido"],
        "tipo": r["Tipo"], "pagador": r["Pagador (origem)"], "valor": float(r["Valor (R$)"]),
        "obs": r.get("Observação") or "",
    } for r in linhas_tabela(fin["Lançamentos"], "Nº")]

    reem = [{
        "n": r["Nº"], "categoria": r["Categoria"], "finalidade": r["Finalidade"],
        "favorecido": r["Favorecido / Serviço"], "data": iso(r["Data e hora"]),
        "forma": r["Forma de pagamento"], "valor": float(r["Valor (R$)"]),
    } for r in linhas_tabela(reemb["Comprovantes"], "Nº")]

    diario = [{
        "n": r["Nº"], "data": iso(r["Data"]), "ambiente": r["Ambiente / Etapa"],
        "efetivo": r.get("Equipe / Efetivo"), "avanco": r.get("Avanço da etapa (%)"),
        "status": r.get("Status"),
    } for r in linhas_tabela(obra["Diário de Obra"], "Nº")]

    # Realizado recalculado a partir dos lançamentos (soma por prefixo do código)
    for item in eap:
        item["real"] = round(sum(l["valor"] for l in lanc if l["cod"].startswith(item["cod"])), 2)
        if abs(item["real"] - item["real_planilha"]) > 0.01:
            print(f"AVISO: realizado de {item['cod']} difere da planilha "
                  f"({item['real']:.2f} x {item['real_planilha']:.2f}). Usando o recalculado.")

    datas = [x["data"][:10] for x in lanc + reem]
    dados = {
        "projeto": "Clínica Sensi · Sala 504 · Rio Mar Trade Center",
        "cidade": "Fortaleza, CE",
        "pasta_origem": Path(pasta).name,
        "periodo": {"inicio": min(datas), "fim": max(datas)},
        "gerado_em": datetime.now().isoformat(timespec="minutes"),
        "mapa_reembolso_eap": MAPA_REEMBOLSO_EAP,
        "eap": eap, "lancamentos": lanc, "reembolsos": reem, "diario": diario,
    }
    Path(saida).parent.mkdir(parents=True, exist_ok=True)
    Path(saida).write_text(json.dumps(dados, ensure_ascii=False, indent=1), encoding="utf-8")
    tot_l = sum(l["valor"] for l in lanc)
    tot_r = sum(r["valor"] for r in reem)
    print(f"OK: {len(eap)} linhas EAP, {len(lanc)} lançamentos (R$ {tot_l:,.2f}), "
          f"{len(reem)} reembolsos (R$ {tot_r:,.2f}), {len(diario)} registros de diário -> {saida}")


if __name__ == "__main__":
    if len(sys.argv) != 3:
        sys.exit(__doc__)
    main(sys.argv[1], sys.argv[2])
